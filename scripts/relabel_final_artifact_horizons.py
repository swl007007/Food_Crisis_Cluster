#!/usr/bin/env python3
"""Relabel paper-facing lag wording to forecasting horizon wording."""

from __future__ import annotations

import argparse
import csv
import json
import sys
from decimal import Decimal, InvalidOperation
from io import StringIO
from pathlib import Path
from typing import NamedTuple

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from paper_horizon_labels import forbidden_paper_lag_terms, replace_paper_horizon_terms


DEFAULT_ROOT = Path("final_artifacts_in_paper_updated")
SUPPORTED_SUFFIXES = {".csv", ".md", ".json", ".xlsx"}
ABLATION_WORKBOOK = "ablation_feature_exclude.xlsx"


class RelabelResult(NamedTuple):
    path: Path
    updated: bool
    remaining_forbidden_terms: list[str]


class TextFileFormat:
    def __init__(self, encoding: str, newline: str) -> None:
        self.encoding = encoding
        self.newline = newline


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


def _detect_text_file_format(path: Path) -> TextFileFormat:
    raw = path.read_bytes()
    crlf_count = raw.count(b"\r\n")
    lf_count = raw.count(b"\n") - crlf_count
    newline = "\r\n" if crlf_count > lf_count else "\n"
    encoding = "utf-8-sig" if raw.startswith(b"\xef\xbb\xbf") else "utf-8"
    return TextFileFormat(encoding=encoding, newline=newline)


def _write_text(path: Path, text: str) -> None:
    file_format = _detect_text_file_format(path)
    with path.open("w", encoding=file_format.encoding, newline=file_format.newline) as handle:
        handle.write(text)



def _csv_rows(text: str) -> list[list[str]]:
    return list(csv.reader(StringIO(text)))


def _numeric_signature(rows: list[list[str]]) -> tuple[tuple[int, int, str], ...]:
    signature: list[tuple[int, int, str]] = []
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            stripped = value.strip()
            if not stripped:
                continue
            try:
                number = Decimal(stripped)
            except InvalidOperation:
                continue
            signature.append((row_index, col_index, str(number.normalize())))
    return tuple(signature)


def _validate_csv_preserved(before: str, after: str, path: Path) -> None:
    before_rows = _csv_rows(before)
    after_rows = _csv_rows(after)
    if len(before_rows) != len(after_rows):
        raise ValueError(f"{path}: CSV row count changed")
    if _numeric_signature(before_rows) != _numeric_signature(after_rows):
        raise ValueError(f"{path}: CSV numeric value signature changed")


def _relabel_text_file(path: Path, dry_run: bool) -> RelabelResult:
    before = _read_text(path)
    after = replace_paper_horizon_terms(before)

    if path.suffix.lower() == ".csv":
        _validate_csv_preserved(before, after, path)
    elif path.suffix.lower() == ".json":
        json.loads(after)

    remaining = forbidden_paper_lag_terms(after)
    if after != before and not dry_run:
        _write_text(path, after)
    return RelabelResult(path=path, updated=after != before, remaining_forbidden_terms=remaining)


def _workbook_strings(path: Path) -> list[str]:
    workbook = load_workbook(path, data_only=False)
    values: list[str] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.data_type != "f":
                    values.append(cell.value)
    workbook.close()
    return values


def _relabel_workbook(path: Path, dry_run: bool) -> RelabelResult:
    workbook = load_workbook(path, data_only=False)
    updated = False
    string_values: list[str] = []

    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str) or cell.data_type == "f":
                        continue
                    replacement = replace_paper_horizon_terms(cell.value)
                    if replacement != cell.value:
                        cell.value = replacement
                        updated = True
                    string_values.append(cell.value)

        if path.name == ABLATION_WORKBOOK and not any("Lag Exclude" in value for value in string_values):
            raise ValueError(f"{path}: expected Lag Exclude row not found")

        remaining = forbidden_paper_lag_terms("\n".join(string_values))
        if updated and not dry_run:
            workbook.save(path)
    finally:
        workbook.close()

    if not dry_run and updated:
        after_values = _workbook_strings(path)
        remaining = forbidden_paper_lag_terms("\n".join(after_values))
    return RelabelResult(path=path, updated=updated, remaining_forbidden_terms=remaining)


def relabel_file(path: Path, dry_run: bool = False) -> RelabelResult:
    """Relabel one supported final-artifact file."""
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        return RelabelResult(path=path, updated=False, remaining_forbidden_terms=[])
    if suffix == ".xlsx":
        return _relabel_workbook(path, dry_run=dry_run)
    return _relabel_text_file(path, dry_run=dry_run)


def iter_supported_files(root: Path) -> list[Path]:
    """Return supported files under root in deterministic order."""
    return sorted(path for path in root.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES)


def relabel_root(root: Path, dry_run: bool = False) -> list[RelabelResult]:
    """Relabel supported files below root."""
    return [relabel_file(path, dry_run=dry_run) for path in iter_supported_files(root)]


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT, help="Final artifact root to relabel")
    parser.add_argument("--dry-run", action="store_true", help="Report updates without writing files")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    root = args.root
    if not root.exists():
        raise FileNotFoundError(f"Artifact root not found: {root}")

    results = relabel_root(root, dry_run=args.dry_run)
    for result in results:
        if result.updated:
            mode = "DRY-RUN UPDATED" if args.dry_run else "UPDATED"
            print(f"{mode} {result.path}")

    remaining = [result for result in results if result.remaining_forbidden_terms]
    if remaining:
        for result in remaining:
            terms = ", ".join(result.remaining_forbidden_terms)
            print(f"FORBIDDEN {result.path}: {terms}", file=sys.stderr)
        return 1

    if not any(result.updated for result in results):
        print(f"No supported files required relabeling under {root}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
