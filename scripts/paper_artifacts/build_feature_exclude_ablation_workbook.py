"""Build the fixed-partition feature-exclude ablation workbook."""

from __future__ import annotations

import argparse
import math
import re
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

try:
    from paper_horizon_labels import HORIZON_MONTHS_BY_SCOPE, label_for_scope
except ModuleNotFoundError:
    from scripts.paper_artifacts.paper_horizon_labels import HORIZON_MONTHS_BY_SCOPE, label_for_scope


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_RUN_ROOT = REPO_ROOT / "main_ablation_exclude_updated_stage3_fixed_partitions"
DEFAULT_MAIN_WORKBOOK = REPO_ROOT / "final_artifacts_in_paper_updated" / "main_month_ind_cont3.xlsx"
DEFAULT_OUTPUT = REPO_ROOT / "final_artifacts_in_paper_updated" / "ablation_feature_exclude.xlsx"

FEATURE_GROUPS = {
    "weather_exclude": "Weather Exclude",
    "agri_exclude": "Agri Exclude",
    "conflict_exclude": "Conflict Exclude",
    "econ_exclude": "Econ Exclude",
    "food_prices_exclude": "Food Prices Exclude",
    "geographic_exclude": "Geographic Exclude",
    "secondary_exclude": "Secondary Exclude",
    "lag_exclude": "Lag Exclude",
}

SCOPE_TO_LAG = {int(scope.removeprefix("fs")): months for scope, months in HORIZON_MONTHS_BY_SCOPE.items()}
LAG_TO_SCOPE = {months: scope for scope, months in SCOPE_TO_LAG.items()}

HEADERS = [
    "",
    "Forecasting horizon",
    "Precision",
    "Recall",
    "F1",
    "precision",
    "recall",
    "F1",
    "F1 Improvement Percentage",
    "F1-Compare with main",
    "F1-compare with main %",
    "F1-compare with baseline",
]


def clean_float(value: Any) -> float | None:
    """Convert numeric values to stable floats for workbook output."""
    if value is None:
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    if math.isnan(result) or math.isinf(result):
        return None
    return round(result, 12)


def parse_forecasting_horizon_months(value: Any) -> int | None:
    """Return horizon months from legacy numeric cells or paper display labels."""
    numeric = clean_float(value)
    if numeric is not None:
        months = int(numeric)
        return months if months in LAG_TO_SCOPE else None

    label = "" if value is None else str(value).strip().lower()
    match = re.search(r"\b(4|8|12)\s*-\s*month\s+(?:horizon|lag)\b", label)
    if match:
        return int(match.group(1))
    return None


def safe_diff(left: Any, right: Any) -> float | None:
    left_num = clean_float(left)
    right_num = clean_float(right)
    if left_num is None or right_num is None:
        return None
    return clean_float(left_num - right_num)


def safe_ratio(numerator: Any, denominator: Any) -> float | None:
    numerator_num = clean_float(numerator)
    denominator_num = clean_float(denominator)
    if numerator_num is None or denominator_num in (None, 0):
        return None
    return clean_float(numerator_num / denominator_num)


def summarize_metrics(path: Path) -> dict[str, float]:
    """Read one Stage 3 metrics_monthly.csv and summarize pooled/partitioned means."""
    if not path.exists():
        raise FileNotFoundError(f"Missing metrics file: {path}")

    df = pd.read_csv(path)
    required = {"model", "precision", "recall", "f1"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"{path} is missing columns: {sorted(missing)}")

    models = set(df["model"].dropna().astype(str))
    if not {"partitioned", "pooled"}.issubset(models):
        raise ValueError(f"{path} must include both partitioned and pooled rows")

    partitioned = df[df["model"] == "partitioned"]
    pooled = df[df["model"] == "pooled"]
    return {
        "Precision": clean_float(partitioned["precision"].mean()),
        "Recall": clean_float(partitioned["recall"].mean()),
        "F1": clean_float(partitioned["f1"].mean()),
        "Pooled precision": clean_float(pooled["precision"].mean()),
        "Pooled recall": clean_float(pooled["recall"].mean()),
        "Pooled F1": clean_float(pooled["f1"].mean()),
        "n_partitioned_months": int(partitioned["test_month"].nunique())
        if "test_month" in partitioned.columns
        else int(len(partitioned)),
        "n_pooled_months": int(pooled["test_month"].nunique())
        if "test_month" in pooled.columns
        else int(len(pooled)),
    }


def _row_label(value: Any, current: str | None) -> str | None:
    if value is None:
        return current
    label = str(value).strip()
    return label or current


def load_main_by_lag(path: Path) -> dict[int, dict[str, float | None]]:
    """Extract GeoRF and FEWSNET reference F1 values from the main workbook."""
    if not path.exists():
        raise FileNotFoundError(f"Missing main workbook: {path}")

    ws = load_workbook(path, data_only=True).active
    refs: dict[int, dict[str, float | None]] = {
        lag: {"partitioned_f1": None, "fewsnet_f1": None}
        for lag in SCOPE_TO_LAG.values()
    }

    current_label: str | None = None
    for row in ws.iter_rows(min_row=3, values_only=True):
        current_label = _row_label(row[0], current_label)
        lag_int = parse_forecasting_horizon_months(row[1])
        if lag_int is None:
            continue
        if lag_int not in refs:
            continue

        label = (current_label or "").lower()
        f1 = clean_float(row[4])
        if label == "georf":
            refs[lag_int]["partitioned_f1"] = f1
        elif "fewsnet" in label or "baseline" in label:
            refs[lag_int]["fewsnet_f1"] = f1

    return refs


def build_ablation_rows(
    run_root: Path,
    main_by_lag: dict[int, dict[str, float | None]],
) -> list[dict[str, Any]]:
    """Build feature-exclude rows from the new fixed-partition run root."""
    rows: list[dict[str, Any]] = []
    for group_key, display in FEATURE_GROUPS.items():
        for scope, lag in SCOPE_TO_LAG.items():
            metrics_path = (
                run_root
                / group_key
                / f"result_partition_k40_compare_GF_fs{scope}"
                / "metrics_monthly.csv"
            )
            summary = summarize_metrics(metrics_path)
            f1 = summary["F1"]
            pooled_f1 = summary["Pooled F1"]
            main_f1 = main_by_lag.get(lag, {}).get("partitioned_f1")
            baseline_f1 = main_by_lag.get(lag, {}).get("fewsnet_f1")
            main_diff = safe_diff(f1, main_f1)

            rows.append(
                {
                    "Feature Group": display,
                    "Forecasting horizon": label_for_scope(f"fs{scope}"),
                    "Precision": summary["Precision"],
                    "Recall": summary["Recall"],
                    "F1": f1,
                    "Pooled precision": summary["Pooled precision"],
                    "Pooled recall": summary["Pooled recall"],
                    "Pooled F1": pooled_f1,
                    "F1 Improvement Percentage": safe_ratio(safe_diff(f1, pooled_f1), pooled_f1),
                    "F1-Compare with main": main_diff,
                    "F1-compare with main %": safe_ratio(main_diff, main_f1),
                    "F1-compare with baseline": safe_diff(f1, baseline_f1),
                }
            )
    return rows


def build_reference_rows(main_workbook: Path) -> list[dict[str, Any]]:
    """Append current GeoRF and FEWSNET reference rows for paper readability."""
    ws = load_workbook(main_workbook, data_only=True).active
    rows: list[dict[str, Any]] = []
    current_label: str | None = None

    for row in ws.iter_rows(min_row=3, values_only=True):
        current_label = _row_label(row[0], current_label)
        if current_label not in {"GeoRF", "FEWSNET (baseline)"}:
            continue
        lag = parse_forecasting_horizon_months(row[1])
        if lag is None:
            continue

        split_f1 = clean_float(row[4])
        pooled_f1 = clean_float(row[7])
        improvement = safe_ratio(safe_diff(split_f1, pooled_f1), pooled_f1)

        rows.append(
            {
                "Feature Group": "Main" if current_label == "GeoRF" else current_label,
                "Forecasting horizon": label_for_scope(f"fs{LAG_TO_SCOPE[lag]}"),
                "Precision": clean_float(row[2]),
                "Recall": clean_float(row[3]),
                "F1": split_f1,
                "Pooled precision": clean_float(row[5]),
                "Pooled recall": clean_float(row[6]),
                "Pooled F1": pooled_f1,
                "F1 Improvement Percentage": improvement,
                "F1-Compare with main": 0 if current_label == "GeoRF" else None,
                "F1-compare with main %": 0 if current_label == "GeoRF" else None,
                "F1-compare with baseline": 0 if current_label == "FEWSNET (baseline)" else None,
            }
        )

    return rows


def write_workbook(rows: list[dict[str, Any]], out_path: Path) -> None:
    """Write rows using the 12-column ablation paper layout."""
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    model_font = Font(bold=True, size=11)
    thin = Border(
        left=Side("thin"),
        right=Side("thin"),
        top=Side("thin"),
        bottom=Side("thin"),
    )
    center = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("C1:E1")
    ws["C1"] = "Split Model"
    ws["C1"].font = header_font
    ws["C1"].alignment = center
    ws["C1"].fill = header_fill

    ws.merge_cells("F1:H1")
    ws["F1"] = "Pooled Model(Non-split)"
    ws["F1"].font = header_font
    ws["F1"].alignment = center
    ws["F1"].fill = header_fill

    for column, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=column, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin

    previous_group: str | None = None
    for row_index, row_data in enumerate(rows, start=3):
        group = row_data["Feature Group"]
        group_value = group if group != previous_group else None
        previous_group = group

        values = [
            group_value,
            row_data["Forecasting horizon"],
            row_data["Precision"],
            row_data["Recall"],
            row_data["F1"],
            row_data["Pooled precision"],
            row_data["Pooled recall"],
            row_data["Pooled F1"],
            row_data["F1 Improvement Percentage"],
            row_data["F1-Compare with main"],
            row_data["F1-compare with main %"],
            row_data["F1-compare with baseline"],
        ]

        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row_index, column=column, value=value)
            cell.border = thin
            if column == 1 and value is not None:
                cell.font = model_font
            if column == 2:
                cell.alignment = center
            if column in {9, 11} and value is not None:
                cell.number_format = "0.00%"

    widths = {
        "A": 24,
        "B": 14,
        "C": 14,
        "D": 14,
        "E": 14,
        "F": 14,
        "G": 14,
        "H": 14,
        "I": 26,
        "J": 20,
        "K": 22,
        "L": 22,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    wb.save(out_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-root", type=Path, default=DEFAULT_RUN_ROOT)
    parser.add_argument("--main-workbook", type=Path, default=DEFAULT_MAIN_WORKBOOK)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--no-reference-rows",
        action="store_true",
        help="Do not append current Main and FEWSNET reference rows.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    main_by_lag = load_main_by_lag(args.main_workbook)
    rows = build_ablation_rows(args.run_root, main_by_lag)
    if not args.no_reference_rows:
        rows.extend(build_reference_rows(args.main_workbook))
    write_workbook(rows, args.out)
    print(f"Saved: {args.out}")
    print(f"Rows written: {len(rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
