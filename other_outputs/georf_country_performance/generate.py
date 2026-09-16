#!/usr/bin/env python3
"""Reproduce the country workbook and figure from frozen main predictions.

Run from the repository root:
    python3 other_outputs/georf_country_performance/generate.py
Requires pandas, numpy, openpyxl, matplotlib and the installed Windows Excel.
"""

from pathlib import Path
import hashlib
import subprocess

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties


OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
ARCHIVE = ROOT / "archived/release_20260624_reproducibility_inputs"
FEWS = ROOT.parents[2] / "1.Source Data/Outcome/FEWSNET_IPC/FEWSNET.csv"
EVALUATOR = ROOT / (
    "archived/release_20260624_nonpaper_pipelines/legacy_misc/app_final/"
    "fewsnet_baseline_evaluation.py"
)
MODELS = ("GeoRF (partitioned)", "Pooled RF", "FEWS NET expert forecasts")
METRICS = ("Precision", "Recall", "F1")
PRED_COLUMNS = ("y_pred_partitioned", "y_pred_pooled", "expert")
KEYS = ["admin_code", "month_start"]
MONTHS = pd.to_datetime([f"{y}-{m:02d}-01" for y in range(2021, 2025)
                         for m in (2, 6, 10)])
SHORT_COVERAGE = {"Afghanistan": 7, "South Sudan": 8, "Burkina Faso": 10,
                  "Somalia": 10, "Burundi": 11, "Sudan": 11, "Yemen": 11}
BOOK = OUT / "country_performance.xlsx"


def sha256(path: Path) -> str:
    """Hash a source without loading it all into memory."""
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def counts_and_scores(truth, prediction) -> tuple[list[int], list[float]]:
    """Positive-class scores; zero denominators match the archived evaluator."""
    y, p = np.asarray(truth), np.asarray(prediction)
    if not (np.isin(y, [0, 1]).all() and np.isin(p, [0, 1]).all()):
        raise ValueError("Expected complete binary labels")
    tp, fp, fn, tn = [int(mask.sum()) for mask in (
        (y == 1) & (p == 1), (y == 0) & (p == 1),
        (y == 1) & (p == 0), (y == 0) & (p == 0),
    )]
    scores = [a / b if b else 0.0 for a, b in
              ((tp, tp + fp), (tp, tp + fn), (2 * tp, 2 * tp + fp + fn))]
    return [tp, fp, fn, tn], scores


def normalize_keys(frame: pd.DataFrame) -> pd.DataFrame:
    """Require integral admin IDs and unique area-month keys."""
    codes = pd.to_numeric(frame["admin_code"], errors="raise")
    if codes.isna().any() or not np.isfinite(codes).all() or (codes % 1 != 0).any():
        raise ValueError("Invalid admin code")
    frame["admin_code"] = codes.astype("int64")
    frame["month_start"] = pd.to_datetime(frame["month_start"], errors="raise")
    if (frame[KEYS].isna().any().any() or frame.duplicated(KEYS).any()
            or frame.month_start.dt.day.ne(1).any()):
        raise ValueError("Invalid or duplicate area-month key")
    return frame.sort_values(KEYS).reset_index(drop=True)


def load_and_verify() -> tuple[dict[int, pd.DataFrame], dict[Path, str]]:
    """Validate source support and reproduce both archived evaluations."""
    sources = [FEWS, EVALUATOR, OUT / "spec.md", Path(__file__).resolve(),
               OUT / "recalculate.ps1"]
    for scope in (1, 2, 3):
        folder = ARCHIVE / f"result_partition_k40_compare_GF_fs{scope}"
        sources.extend([folder / "predictions_monthly.csv", folder / "metrics_monthly.csv"])
    sources.extend(ARCHIVE / "fewsnet_baseline_results" /
                   f"fewsnet_baseline_results_fs{s}.csv" for s in (1, 2))
    hashes = {p: sha256(p) for p in sources}
    raw = pd.read_csv(FEWS)
    for name in ("admin_code", "year", "month"):
        raw[name] = pd.to_numeric(raw[name], errors="raise")
    invalid = raw[["admin_code", "year", "month"]].isna().any(axis=1)
    if invalid.sum() != 1 or not raw.loc[invalid, ["admin_code", "year", "month"]].isna().all().all():
        raise ValueError("Unexpected invalid source rows")
    raw = raw.loc[~invalid].copy()
    if ((raw[["year", "month"]] % 1 != 0).any().any()
            or not raw.month.between(1, 12).all()):
        raise ValueError("Invalid source date")
    raw["month_start"] = pd.to_datetime(dict(year=raw.year.astype(int),
                                            month=raw.month.astype(int), day=1))
    raw = normalize_keys(raw)
    raw["country"] = raw.country.str.strip()
    if raw.country.isna().any() or raw.country.eq("").any():
        raise ValueError("Missing country mapping")
    raw["source_truth"] = raw.fews_ipc.ge(3).astype(int)
    raw["quarter"] = (raw.month.astype(int) - 1) // 3 + 1
    # Deliberately preserve archived missing-phase -> 0 BEFORE record shifting.
    for scope, phase in ((1, "fews_proj_near"), (2, "fews_proj_med")):
        binary = raw[phase].ge(3).astype(int)
        raw[f"expert_{scope}"] = binary.groupby(raw.admin_code).shift(scope * 4)
        reference = pd.read_csv(ARCHIVE / "fewsnet_baseline_results" /
                                f"fewsnet_baseline_results_fs{scope}.csv")
        if len(reference) != 39 or reference.duplicated(["year", "quarter"]).any():
            raise ValueError("Unexpected archived baseline support")
        for row in reference.to_dict("records"):
            quarter = raw.loc[raw.year.eq(row["year"]) & raw.quarter.eq(row["quarter"])]
            quarter = quarter.dropna(subset=[f"expert_{scope}"])
            counts, scores = counts_and_scores(quarter.source_truth, quarter[f"expert_{scope}"])
            np.testing.assert_allclose(scores, [row[f"{m.lower()}(1)"] for m in METRICS],
                                       rtol=0, atol=1e-12)
            assert counts[0] + counts[2] == row["num_samples(1)"]
        print(f"fs{scope}: all 39 archived expert quarters reproduced", flush=True)

    frames = {}
    for scope in (1, 2, 3):
        folder = ARCHIVE / f"result_partition_k40_compare_GF_fs{scope}"
        frame = normalize_keys(pd.read_csv(folder / "predictions_monthly.csv").rename(
            columns={"FEWSNET_admin_code": "admin_code"}))
        if len(frame) != 62189 or set(frame.month_start) != set(MONTHS):
            raise ValueError("Unexpected main prediction support")
        if scope > 1:
            pd.testing.assert_frame_equal(frame[KEYS + ["y_true"]], frames[4][KEYS + ["y_true"]])
        frame = frame.merge(raw[KEYS + ["country", "fews_ipc", "source_truth", "expert_1", "expert_2"]],
                            on=KEYS, how="left", validate="one_to_one", indicator=True)
        if (not frame._merge.eq("both").all() or frame.country.isna().any()
                or frame.fews_ipc.isna().any() or not frame.y_true.eq(frame.source_truth).all()):
            raise ValueError("Missing source match or mismatched observed truth")
        if scope < 3:
            frame["expert"] = frame[f"expert_{scope}"]
            if frame.expert.isna().any():
                raise ValueError("Missing expert label on common support")
        reference = pd.read_csv(folder / "metrics_monthly.csv")
        if len(reference) != 24 or reference.duplicated(["test_month", "model"]).any():
            raise ValueError("Unexpected archived model metric support")
        for model in ("partitioned", "pooled"):
            for month, group in frame.groupby("month_start"):
                counts, scores = counts_and_scores(group.y_true, group[f"y_pred_{model}"])
                match = reference.loc[reference.model.eq(model) & reference.test_month.eq(month.strftime("%Y-%m"))]
                assert len(match) == 1
                row = match.iloc[0]
                np.testing.assert_array_equal(counts, row[["tp", "fp", "fn", "tn"]].to_numpy())
                np.testing.assert_allclose(scores, row[[m.lower() for m in METRICS]].to_numpy(dtype=float),
                                           rtol=0, atol=1e-12)
                assert row["n"] == len(group)
        coverage = frame.groupby("country").month_start.nunique()
        if len(coverage) != 22 or any(n != SHORT_COVERAGE.get(c, 12) for c, n in coverage.items()):
            raise ValueError("Unexpected country/month coverage")
        frames[scope * 4] = frame
        print(f"fs{scope}: 24 model rows reproduced; 62,189 matched rows, 22 countries", flush=True)
    return frames, hashes


def monthly_statistics(frames: dict[int, pd.DataFrame]) -> pd.DataFrame:
    """Observed country-months only, with common-support confusion counts."""
    rows = []
    for horizon, frame in frames.items():
        for (country, month), group in frame.groupby(["country", "month_start"]):
            row = [country, horizon, month.strftime("%Y-%m"), len(group)]
            for column in PRED_COLUMNS:
                if column == "expert" and horizon == 12:
                    row.extend([np.nan] * 7)
                else:
                    counts, scores = counts_and_scores(group.y_true, group[column])
                    row.extend(counts + scores)
            rows.append(row)
    columns = ["country", "horizon", "month", "n"]
    columns.extend(f"{model}_{stat}" for model in range(3)
                   for stat in ("tp", "fp", "fn", "tn", "Precision", "Recall", "F1"))
    return pd.DataFrame(rows, columns=columns).sort_values(["country", "horizon", "month"]).reset_index(drop=True)


def write_workbook(monthly: pd.DataFrame, hashes: dict[Path, str]) -> None:
    """Store source counts, formula ratios and formula equal-month averages."""
    book = Workbook()
    main = book.active
    main.title = "Country performance"
    main.merge_cells("A1:M1")
    main["A1"] = "Country-level forecast performance | IPC Phase 3+"
    main.merge_cells("A2:M2")
    main["A2"] = "2021–2024 · February / June / October · equal weight per observed month"
    headers = ["Country", "Horizon (months)", "N area-months", "Observed months"] + list(METRICS) * 3
    for col, title in enumerate(headers, 1):
        main.cell(4, col, title)
    for model, name in enumerate(MODELS):
        start = 5 + model * 3
        main.merge_cells(start_row=3, start_column=start, end_row=3, end_column=start + 2)
        main.cell(3, start, name)
    sheet = book.create_sheet("Monthly counts")
    sheet.append(["Country", "Horizon (months)", "Evaluation month", "N area-months"] +
                 [f"{model} | {stat}" for model in MODELS for stat in ("TP", "FP", "FN", "TN", *METRICS)])
    for idx, record in enumerate(monthly.itertuples(index=False, name=None), 2):
        sheet.append([*record[:3], f"=SUM(E{idx}:H{idx})"])
        for model in range(3):
            start = 5 + model * 7
            if record[1] == 12 and model == 2:
                for col in range(start, start + 7):
                    sheet.cell(idx, col, "N/A")
                continue
            for k in range(4):
                sheet.cell(idx, start + k, int(record[4 + model * 7 + k]))
            tp, fp, fn = [f"{get_column_letter(start + k)}{idx}" for k in range(3)]
            for k, formula in enumerate((f"{tp}/({tp}+{fp})", f"{tp}/({tp}+{fn})",
                                         f"2*{tp}/(2*{tp}+{fp}+{fn})")):
                sheet.cell(idx, start + 4 + k, f"=IFERROR({formula},0)")
    for idx, ((country, horizon), group) in enumerate(monthly.groupby(["country", "horizon"], sort=True), 5):
        first, last = int(group.index.min()) + 2, int(group.index.max()) + 2
        main.cell(idx, 1, country)
        main.cell(idx, 2, int(horizon))
        main.cell(idx, 3, f"=SUM('Monthly counts'!D{first}:D{last})")
        main.cell(idx, 4, f"=COUNT('Monthly counts'!D{first}:D{last})")
        for model in range(3):
            for metric in range(3):
                col = get_column_letter(9 + model * 7 + metric)
                value = "N/A" if horizon == 12 and model == 2 else f"=AVERAGE('Monthly counts'!{col}{first}:{col}{last})"
                main.cell(idx, 5 + model * 3 + metric, value)

    notes = book.create_sheet("Notes and sources")
    notes.append(["Item", "Definition / source", "SHA-256"])
    for item, detail in [
        ("Reproduce", "python3 other_outputs/georf_country_performance/generate.py (from repository root; WSL with installed Windows Excel)"),
        ("Target", "IPC Phase 3+; raw fews_ipc >= 3 agrees exactly with saved binary y_true on all selected rows."),
        ("Model predictions", "Use saved y_pred_partitioned / y_pred_pooled exactly. No probability thresholding, fitting or threshold selection."),
        ("Expert convention", "Original evaluator: near/medium phase >= 3; missing phase becomes 0; sort complete admin history by year/month; per-admin record shift(4)/shift(8) BEFORE evaluation filtering. These are record shifts, not calendar joins."),
        ("Support", "62,189 identical area-month keys and truth per horizon; 22 countries. All models at 4/8 months share these rows. Main sample need not match full archived expert support."),
        ("Historical invalid row", "Drop the single source row with null admin_code/year/month before date conversion (country field: System.IO.MemoryStream). No other source rows dropped before shifting."),
        ("Monthly formulas", "TP/(TP+FP), TP/(TP+FN), 2TP/(2TP+FP+FN). A zero denominator yields 0 on observed country-months. N is TP+FP+FN+TN."),
        ("Country formulas", "Arithmetic average of observed monthly scores; monthly F1 is calculated before averaging. N sums monthly area counts; observed months counts existing monthly rows."),
        ("Coverage", "Afghanistan 7; South Sudan 8; Burkina Faso/Somalia 10; Burundi/Sudan/Yemen 11; remaining countries 12. Missing country-months are omitted, never scored zero."),
        ("Unavailable", "12-month FEWS NET forecasts: literal N/A, no 8-month proxy. This differs from absent country-months, which have no row."),
        ("Verification", "All 72 archived model month/model rows and all 78 archived expert quarter/scope rows reproduced within 1e-12, including original positive support. All workbook formula caches independently checked after Excel recalculation."),
        ("Figure", "Descriptive 3 × 3 comparison; country order alphabetical; no lines between categories, ranks, intervals or significance tests. Points use verified cached country-sheet values. DR Congo abbreviates Democratic Republic of the Congo."),
        ("Source paths", "Paths below are relative to repository root except the external FEWS NET source. Sources and generator are SHA-256 checked again after generation."),
        ("Runtime", f"pandas {pd.__version__}; numpy {np.__version__}; matplotlib {matplotlib.__version__}; workbook recalculated in installed Microsoft Excel"),
    ]:
        notes.append([item, detail])
    for path, digest in hashes.items():
        notes.append(["Source", str(path.relative_to(ROOT)) if path.is_relative_to(ROOT) else str(path), digest])

    dark = "203B4D"
    group_colors = ("DDECF4", "E9EBEF", "F7E7D0")
    for ws, header_row in ((main, 4), (sheet, 1), (notes, 1)):
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = f"C{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
        for row in ws:
            for cell in row:
                cell.font = Font(name="Arial", size=10, color="202C35")
                cell.alignment = Alignment(vertical="center", horizontal="left" if cell.column == 1 else "center")
                if cell.row == header_row:
                    cell.fill = PatternFill("solid", fgColor=dark)
                    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                    cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
                elif cell.row > header_row and cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F2F5F7")
        ws.row_dimensions[header_row].height = 34 if ws is main else 48
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = f"1:{header_row}"
    main["A1"].font = Font(name="Arial", size=17, bold=True, color=dark)
    main["A2"].alignment = Alignment(horizontal="left")
    main.row_dimensions[1].height = 30
    main.row_dimensions[2].height = 24
    main.row_dimensions[3].height = 30
    main.column_dimensions["A"].width = 36
    for col in range(2, 14):
        main.column_dimensions[get_column_letter(col)].width = 15
    for model in range(3):
        for col in range(5 + model * 3, 8 + model * 3):
            main.cell(3, col).fill = PatternFill("solid", fgColor=group_colors[model])
        main.cell(3, 5 + model * 3).font = Font(name="Arial", size=11, bold=True, color=dark)
    for row in main.iter_rows(min_row=5):
        main.row_dimensions[row[0].row].height = 21
        for cell in row[4:]:
            cell.number_format = "0.000"
        row[2].number_format = "#,##0"
    sheet.column_dimensions["A"].width = 36
    for col in range(2, 26):
        sheet.column_dimensions[get_column_letter(col)].width = 16
    for row in sheet.iter_rows(min_row=2):
        for model in range(3):
            for cell in row[8 + model * 7:11 + model * 7]:
                cell.number_format = "0.000"
    for col, width in (("A", 25), ("B", 116), ("C", 72)):
        notes.column_dimensions[col].width = width
    for row in notes.iter_rows(min_row=2):
        notes.row_dimensions[row[0].row].height = 48
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    book.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True)
    book.save(BOOK)


def verify_workbook(monthly: pd.DataFrame) -> pd.DataFrame:
    """Independently compare every cached numeric result against numpy/pandas."""
    cached = load_workbook(BOOK, data_only=True)
    formulas = load_workbook(BOOK, data_only=False)
    formula_count = 0
    for ws in formulas:
        for row in ws:
            for cell in row:
                value = cached[ws.title][cell.coordinate]
                assert value.data_type != "e", (ws.title, cell.coordinate, value.value)
                if cell.data_type == "f":
                    formula_count += 1
                    assert isinstance(value.value, (int, float)) and np.isfinite(value.value)
    sheet = cached["Monthly counts"]
    assert sheet.max_row == len(monthly) + 1
    for idx, record in enumerate(monthly.itertuples(index=False, name=None), 2):
        for col, expected in enumerate(record, 1):
            value = sheet.cell(idx, col).value
            if isinstance(expected, str):
                assert value == expected
            elif pd.isna(expected):
                assert value == "N/A"
            else:
                np.testing.assert_allclose(value, expected, rtol=0, atol=1e-12)
    main = cached["Country performance"]
    assert main.max_row == 70
    results = []
    for idx, ((country, horizon), group) in enumerate(monthly.groupby(["country", "horizon"], sort=True), 5):
        values = [cell.value for cell in main[idx]]
        assert values[:4] == [country, horizon, int(group.n.sum()), len(group)]
        scores = group[[f"{m}_{s}" for m in range(3) for s in METRICS]].mean().to_numpy()
        for value, expected in zip(values[4:], scores):
            if np.isnan(expected):
                assert value == "N/A"
            else:
                np.testing.assert_allclose(value, expected, rtol=0, atol=1e-12)
        results.append(values)
    cached.close()
    formulas.close()
    print(f"Workbook: 66 country/horizon rows, {len(monthly)} monthly rows; {formula_count} cached formulas verified; zero Excel errors", flush=True)
    return pd.DataFrame(results, columns=["country", "horizon", "n", "months"] +
                        [f"{m}_{s}" for m in range(3) for s in METRICS])


def draw_figure(summary: pd.DataFrame) -> None:
    """Plot the verified workbook scores as offset country-category points."""
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10,
                         "pdf.fonttype": 42, "axes.spines.top": False,
                         "axes.spines.right": False, "axes.linewidth": 0.7})
    countries = sorted(summary.country.unique())
    labels = ["DR Congo" if c == "Democratic Republic of the Congo" else c for c in countries]
    colors, markers, offsets = ("#0072B2", "#6B6B6B", "#D58D00"), ("o", "s", "^"), (-0.22, 0, 0.22)
    fig, axes = plt.subplots(3, 3, figsize=(18, 13), sharex=True, sharey=True)
    fig.subplots_adjust(left=0.065, right=0.99, bottom=0.205, top=0.90, hspace=0.15, wspace=0.10)
    for row, horizon in enumerate((4, 8, 12)):
        data = summary.loc[summary.horizon.eq(horizon)].set_index("country").loc[countries]
        for col, metric in enumerate(METRICS):
            ax = axes[row, col]
            for model in range(3 if horizon < 12 else 2):
                ax.scatter(np.arange(22) + offsets[model], data[f"{model}_{metric}"].to_numpy(dtype=float),
                           s=23, marker=markers[model], color=colors[model], linewidths=0.4,
                           edgecolors="white", zorder=3, clip_on=False)
            ax.set_ylim(0, 1)
            ax.set_xlim(-0.7, 21.7)
            ax.set_yticks(np.arange(0, 1.01, 0.2))
            ax.set_axisbelow(True)
            ax.grid(axis="y", color="#DDE1E4", linewidth=0.6)
            ax.tick_params(axis="x", length=2, pad=4)
            ax.tick_params(axis="y", labelsize=10)
            ax.set_xticks(np.arange(22), labels, rotation=65, ha="right", fontsize=9)
            ax.text(0.01, 1.04, chr(97 + row * 3 + col), transform=ax.transAxes, fontweight="bold", fontsize=12)
            if row == 0:
                ax.set_title(metric, fontsize=14, pad=15)
            if col == 0:
                ax.set_ylabel(f"{horizon}-month horizon\nScore", fontsize=12, labelpad=13)
    fig.suptitle("Country-level performance of food-crisis forecasts", y=0.975, fontsize=19)
    handles = [Line2D([], [], color=color, marker=marker, linestyle="None", markersize=7, label=name)
               for color, marker, name in zip(colors, markers, MODELS)]
    fig.legend(handles=handles, loc="upper center", bbox_to_anchor=(0.53, 0.955), ncol=3,
               frameon=False, fontsize=12, handletextpad=0.5, columnspacing=2.5)
    fig.text(0.065, 0.061, "IPC Phase 3+ · Scores average observed country-months equally (7–12 months per country), Feb/Jun/Oct 2021–2024.", fontsize=10)
    fig.text(0.065, 0.039, "Identical area-month support within each horizon; 12-month FEWS NET expert forecasts are unavailable. Source values and coverage: country_performance.xlsx.", fontsize=10)
    fig.savefig(OUT / "country_performance_3x3.png", dpi=300, facecolor="white")
    fig.savefig(OUT / "country_performance_3x3.pdf", facecolor="white")
    plt.close(fig)


def main() -> None:
    """Generate, recalculate, verify, and render only this directory's outputs."""
    frames, hashes = load_and_verify()
    monthly = monthly_statistics(frames)
    write_workbook(monthly, hashes)
    print("Recalculating the new workbook in a separate Excel instance...", flush=True)
    windows_paths = [subprocess.check_output(["wslpath", "-w", str(p)], text=True).strip()
                     for p in (OUT / "recalculate.ps1", BOOK)]
    subprocess.run(["/mnt/c/Windows/System32/WindowsPowerShell/v1.0/powershell.exe",
                    "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                    "-File", windows_paths[0], "-WorkbookPath", windows_paths[1]], check=True)
    summary = verify_workbook(monthly)
    draw_figure(summary)
    for path, before in hashes.items():
        assert sha256(path) == before, f"Source changed during generation: {path}"
    print("All source hashes unchanged. Workbook, 300-dpi PNG and vector PDF generated.", flush=True)


if __name__ == "__main__":
    main()
