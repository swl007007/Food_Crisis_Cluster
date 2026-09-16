"""Map cached, equal-observed-month country F1 differences; never refit models.

The five quantitative panels show where partitioned GeoRF improves on each
comparator in 2021-2024. The unavailable 12-month expert panel has no axes.
Run with the existing Windows Python 3.12 environment; all inputs are local.
"""

from hashlib import sha256
from math import ceil
from pathlib import Path

import geopandas as gpd
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.colors import TwoSlopeNorm
from openpyxl import load_workbook
from shapely.geometry import box


OUTPUT_DIR = Path(__file__).resolve().parent
ROOT = OUTPUT_DIR.parents[1]
WORKBOOK = OUTPUT_DIR / "country_performance.xlsx"
FEWS_BOUNDARIES = (
    ROOT.parents[2] / "1.Source Data" / "Outcome" / "FEWSNET_IPC"
    / "FEWS NET Admin Boundaries" / "FEWS_Admin_LZ_v3.shp"
)
CONTEXT_BOUNDARIES = (
    Path.home() / ".local/share/cartopy/shapefiles/natural_earth/cultural"
    / "ne_110m_admin_0_countries.shp"
)
LATAM_COUNTRIES = ("Guatemala", "Haiti")
# Same study footprint, projection, simplification and inset as the map reference.
AFRICA_MIN_X_M = -2_226_000.0
SIMPLIFY_M = 5_000
INSET_POSITION = [0.01, 0.01, 0.30, 0.28]


def load_country_differences() -> pd.DataFrame:
    """Subtract full-precision cached F1 cells without another aggregation."""
    workbook = load_workbook(WORKBOOK, data_only=True, read_only=True)
    sheet = workbook["Country performance"]
    rows = list(sheet.iter_rows(min_row=5, max_row=70, values_only=True))
    workbook.close()
    data = pd.DataFrame(rows, columns=[
        "country", "horizon", "n", "months", "gp", "gr", "georf",
        "pp", "pr", "pooled", "ep", "er", "expert",
    ])
    assert len(data) == 66 and data.country.nunique() == 22
    assert not data.duplicated(["country", "horizon"]).any()
    assert set(data.horizon) == {4, 8, 12}
    assert data.groupby("horizon").size().eq(22).all()
    assert data.groupby("horizon").n.sum().eq(62189).all()
    assert data.groupby("country").n.nunique().eq(1).all()
    unavailable = data.horizon.eq(12)
    assert data.loc[unavailable, "expert"].eq("N/A").all()
    data.loc[unavailable, "expert"] = np.nan
    for column in ("georf", "pooled", "expert"):
        data[column] = pd.to_numeric(data[column], errors="raise")
        scored = data.loc[~unavailable if column == "expert" else slice(None), column]
        assert np.isfinite(scored).all() and scored.between(0, 1).all()
    data["delta_expert"] = data.georf - data.expert
    data["delta_pooled"] = data.georf - data.pooled
    assert data[["delta_expert", "delta_pooled"]].count().sum() == 110
    return data


def load_country_geometry(countries: set[str]) -> tuple:
    """Dissolve study boundaries by country; Natural Earth supplies context only."""
    study = gpd.read_file(FEWS_BOUNDARIES)[["ADMIN0", "geometry"]]
    context = gpd.read_file(CONTEXT_BOUNDARIES)
    assert study.crs.to_epsg() == context.crs.to_epsg() == 4326
    assert len(study) == 5718 and set(study.ADMIN0) == countries
    assert study.geometry.notna().all() and not study.geometry.is_empty.any()
    projected = study.to_crs(3857)
    main_bounds = projected.loc[
        projected.geometry.centroid.x >= AFRICA_MIN_X_M
    ].total_bounds
    latam_bounds = study.loc[study.ADMIN0.isin(LATAM_COUNTRIES)].total_bounds
    latam_pad = (latam_bounds[2:] - latam_bounds[:2]) * [0.05, 0.12]
    latam_bounds = np.r_[latam_bounds[:2] - latam_pad,
                         latam_bounds[2:] + latam_pad]
    invalid = ~study.geometry.is_valid
    study.loc[invalid, "geometry"] = study.loc[invalid].geometry.make_valid()
    country = study.dissolve(by="ADMIN0")
    assert len(country) == 22 and set(country.index) == countries
    assert country.geometry.is_valid.all() and not country.geometry.is_empty.any()
    country = country.to_crs(3857)
    # Projection can turn touching polygon edges into tiny self-intersections.
    country.geometry = country.geometry.make_valid()
    country.geometry = country.geometry.simplify(SIMPLIFY_M, preserve_topology=True)
    country.geometry = country.geometry.make_valid()
    assert country.geometry.is_valid.all() and not country.geometry.is_empty.any()
    main = country.loc[~country.index.isin(LATAM_COUNTRIES)]
    latam = country.loc[list(LATAM_COUNTRIES)].to_crs(4326)
    assert len(main) == 20 and len(latam) == 2
    assert main.intersects(box(*main_bounds)).all()
    assert latam.intersects(box(*latam_bounds)).all()
    print(f"Geometry: {len(country)} countries; repaired {invalid.sum()} polygons")
    return main, latam, context.to_crs(3857), context, main_bounds, latam_bounds


def draw_country_map(ax, geometry, context, values, bounds, norm) -> None:
    """Draw one country color above neutral context, with fixed shared limits."""
    minx, miny, maxx, maxy = bounds
    context.cx[minx:maxx, miny:maxy].plot(
        ax=ax, color="#eeeeec", edgecolor="#a5a5a5", linewidth=0.30, zorder=1,
    )
    layer = geometry.assign(delta=geometry.index.map(values))
    assert layer.delta.notna().all()
    # Geometry repair can retain collapsed edges as lines, not country areas.
    layer = layer.explode(index_parts=False)
    layer = layer.loc[layer.geom_type.isin(("Polygon", "MultiPolygon"))]
    assert set(layer.index) == set(geometry.index)
    layer.plot(
        ax=ax, column="delta", cmap="RdBu", norm=norm,
        edgecolor="none", linewidth=0, zorder=2,
    )
    ax.set_xlim(minx, maxx)
    ax.set_ylim(miny, maxy)
    ax.set_xticks([])
    ax.set_yticks([])


def main() -> None:
    """Validate inputs and export only the new country F1 map PNG and PDF."""
    workbook_hash = sha256(WORKBOOK.read_bytes()).hexdigest()
    data = load_country_differences()
    main_geo, latam, main_context, latam_context, bounds, inset_bounds = (
        load_country_geometry(set(data.country))
    )
    deltas = data[["delta_expert", "delta_pooled"]].to_numpy(dtype=float)
    scale = ceil(float(np.nanmax(np.abs(deltas))) * 10) / 10
    assert scale > 0 and np.nanmax(np.abs(deltas)) <= scale
    norm = TwoSlopeNorm(vmin=-scale, vcenter=0, vmax=scale)
    plt.rcParams.update({
        "font.family": "sans-serif", "font.sans-serif": ["Arial", "DejaVu Sans"],
        "font.size": 11, "pdf.fonttype": 42, "ps.fonttype": 42,
    })
    fig = plt.figure(figsize=(14, 17.5), facecolor="white")
    grid = fig.add_gridspec(3, 2, left=0.065, right=0.99, bottom=0.115,
                           top=0.925, hspace=0.045, wspace=0.055)
    axes, insets = [], []
    letter = iter("abcde")
    for row, horizon in enumerate((4, 8, 12)):
        row_box = grid[row, 0].get_position(fig)
        fig.text(0.027, (row_box.y0 + row_box.y1) / 2,
                 f"{horizon}-month horizon", rotation=90, ha="center", va="center",
                 fontsize=13, fontweight="bold")
        for col, comparison in enumerate(("expert", "pooled")):
            if (row, col) == (2, 0):
                continue  # No axes, label, frame, inset or other artist in this cell.
            ax = fig.add_subplot(grid[row, col])
            values = data.loc[data.horizon.eq(horizon)].set_index("country")[
                f"delta_{comparison}"
            ]
            assert values.count() == 22
            draw_country_map(ax, main_geo, main_context, values, bounds, norm)
            ax.set_axis_off()
            ax.text(0, 1.015, next(letter), transform=ax.transAxes,
                    fontsize=15, fontweight="bold", ha="left", va="bottom")
            inset = ax.inset_axes(INSET_POSITION)
            draw_country_map(inset, latam, latam_context, values, inset_bounds, norm)
            for spine in inset.spines.values():
                spine.set_color("0.35")
                spine.set_linewidth(0.9)
            inset.set_title("Latin America", fontsize=8, pad=1.5)
            inset.set_facecolor("white")
            axes.append(ax)
            insets.append(inset)
    for col, title in enumerate(("FEWS NET expert forecasts", "Pooled RF")):
        column_box = grid[0, col].get_position(fig)
        fig.text((column_box.x0 + column_box.x1) / 2, 0.965,
                 f"GeoRF (partitioned) −\n{title}", ha="center", va="top",
                 fontsize=14, fontweight="bold", linespacing=1.45)
    colorbar = fig.colorbar(
        plt.cm.ScalarMappable(norm=norm, cmap="RdBu"),
        cax=fig.add_axes([0.31, 0.073, 0.40, 0.012]), orientation="horizontal",
        ticks=np.linspace(-scale, scale, 5),
    )
    colorbar.set_label("ΔF1 (GeoRF − comparator)", fontsize=12, labelpad=6)
    colorbar.outline.set_linewidth(0.5)
    fig.text(0.5, 0.026, "Blue (positive) favours GeoRF; red (negative) favours comparator.",
             ha="center", fontsize=10)
    fig.text(0.5, 0.012,
             "F1 averaged equally over observed country months, 2021–2024 (Feb/Jun/Oct). "
             "12-month FEWS NET forecasts unavailable.", ha="center", fontsize=9)
    assert len(axes) == len(insets) == 5
    assert all(np.allclose(ax.get_xlim(), bounds[[0, 2]]) for ax in axes)
    assert all(np.allclose(ax.get_ylim(), bounds[[1, 3]]) for ax in axes)
    assert all(np.allclose(ax.get_xlim(), inset_bounds[[0, 2]]) for ax in insets)
    assert all(np.allclose(ax.get_ylim(), inset_bounds[[1, 3]]) for ax in insets)
    description = (
        f"Workbook SHA256={workbook_hash}; geometry={FEWS_BOUNDARIES.name}; "
        f"context={CONTEXT_BOUNDARIES.name}; ADMIN0 country dissolve; EPSG3857 main; "
        f"EPSG4326 inset; simplification={SIMPLIFY_M}m; symmetric limits={-scale},{scale}"
    )
    output = OUTPUT_DIR / "country_f1_differences_3x2"
    fig.savefig(output.with_suffix(".png"), dpi=300,
                metadata={"Description": description})
    fig.savefig(output.with_suffix(".pdf"),
                metadata={"Title": "Country-level F1 differences", "Subject": description})
    plt.close(fig)
    assert sha256(WORKBOOK.read_bytes()).hexdigest() == workbook_hash
    print(f"Verified 66 rows, 110 differences; range [{np.nanmin(deltas):.9f}, "
          f"{np.nanmax(deltas):.9f}]; shared scale [{-scale}, {scale}]")
    print(f"Unchanged workbook SHA256: {workbook_hash}")
    print(f"Saved {output.name}.png and .pdf")


if __name__ == "__main__":
    main()
