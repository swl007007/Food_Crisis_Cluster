from pathlib import Path
import tempfile
import unittest

import pandas as pd
from openpyxl import Workbook
from openpyxl import load_workbook

from scripts.build_feature_exclude_ablation_workbook import (
    FEATURE_GROUPS,
    build_ablation_rows,
    build_reference_rows,
    load_main_by_lag,
    write_workbook,
)


def _write_metrics(path: Path, partitioned_f1: float, pooled_f1: float) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(
        [
            {
                "test_month": "2021-02",
                "model": "partitioned",
                "precision": 0.7,
                "recall": 0.5,
                "f1": partitioned_f1,
            },
            {
                "test_month": "2021-02",
                "model": "pooled",
                "precision": 0.6,
                "recall": 0.4,
                "f1": pooled_f1,
            },
            {
                "test_month": "2021-03",
                "model": "partitioned",
                "precision": 0.9,
                "recall": 0.7,
                "f1": partitioned_f1 + 0.2,
            },
            {
                "test_month": "2021-03",
                "model": "pooled",
                "precision": 0.8,
                "recall": 0.6,
                "f1": pooled_f1 + 0.2,
            },
        ]
    ).to_csv(path, index=False)


class FeatureExcludeAblationWorkbookTests(unittest.TestCase):
    def test_main_workbook_readers_accept_horizon_display_labels(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "main_month_ind_cont3.xlsx"
            wb = Workbook()
            ws = wb.active
            rows = [
                ("GeoRF", "4-month horizon", 0.7, 0.5, 0.60, 0.6, 0.4, 0.50),
                (None, "8-month horizon", 0.8, 0.6, 0.70, 0.7, 0.5, 0.60),
                (None, "12-month horizon", 0.9, 0.7, 0.80, 0.8, 0.6, 0.70),
                ("FEWSNET (baseline)", "4-month horizon", 0.4, 0.3, 0.35, None, None, None),
                (None, "8-month horizon", 0.5, 0.4, 0.45, None, None, None),
                (None, "12-month horizon", 0.6, 0.5, 0.55, None, None, None),
            ]
            for row in rows:
                ws.append(row)
            ws.insert_rows(1, amount=2)
            wb.save(path)

            refs = load_main_by_lag(path)
            reference_rows = build_reference_rows(path)

        self.assertEqual(refs[4]["partitioned_f1"], 0.60)
        self.assertEqual(refs[8]["partitioned_f1"], 0.70)
        self.assertEqual(refs[12]["partitioned_f1"], 0.80)
        self.assertEqual(refs[4]["fewsnet_f1"], 0.35)
        self.assertEqual(refs[8]["fewsnet_f1"], 0.45)
        self.assertEqual(refs[12]["fewsnet_f1"], 0.55)
        self.assertEqual([row["Forecasting horizon"] for row in reference_rows[:3]], [
            "4-month horizon",
            "8-month horizon",
            "12-month horizon",
        ])

    def test_build_ablation_rows_uses_new_run_root(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            run_root = Path(tmp) / "runs"
            for group in FEATURE_GROUPS:
                for scope in (1, 2, 3):
                    _write_metrics(
                        run_root
                        / group
                        / f"result_partition_k40_compare_GF_fs{scope}"
                        / "metrics_monthly.csv",
                        0.50 + scope / 10,
                        0.40 + scope / 10,
                    )

            main_by_lag = {
                4: {"partitioned_f1": 0.80, "fewsnet_f1": 0.30},
                8: {"partitioned_f1": 0.70, "fewsnet_f1": 0.20},
                12: {"partitioned_f1": 0.60, "fewsnet_f1": 0.10},
            }
            rows = build_ablation_rows(run_root, main_by_lag)

        self.assertEqual(len(rows), 24)
        self.assertEqual(rows[0]["Feature Group"], "Weather Exclude")
        self.assertEqual(rows[0]["Forecasting horizon"], "4-month horizon")
        self.assertEqual(rows[0]["F1"], 0.7)
        self.assertEqual(rows[0]["Pooled F1"], 0.6)
        self.assertAlmostEqual(
            rows[0]["F1 Improvement Percentage"], (0.7 - 0.6) / 0.6
        )
        self.assertEqual(rows[0]["F1-Compare with main"], -0.1)
        self.assertEqual(rows[0]["F1-compare with baseline"], 0.4)
        self.assertEqual(rows[-6]["Feature Group"], "Secondary Exclude")
        self.assertEqual(rows[-3]["Feature Group"], "Lag Exclude")

    def test_write_workbook_preserves_paper_shape(self) -> None:
        rows = []
        for group in FEATURE_GROUPS:
            for lag in (4, 8, 12):
                rows.append(
                    {
                        "Feature Group": group.replace("_", " ").title(),
                        "Forecasting horizon": f"{lag}-month horizon",
                        "Precision": 0.7,
                        "Recall": 0.5,
                        "F1": 0.6,
                        "Pooled precision": 0.6,
                        "Pooled recall": 0.4,
                        "Pooled F1": 0.5,
                        "F1 Improvement Percentage": 0.2,
                        "F1-Compare with main": -0.1,
                        "F1-compare with main %": -0.1428571429,
                        "F1-compare with baseline": 0.3,
                    }
                )

        with tempfile.TemporaryDirectory() as tmp:
            out_path = Path(tmp) / "ablation_feature_exclude.xlsx"
            write_workbook(rows, out_path)
            ws = load_workbook(out_path, data_only=False).active

        self.assertEqual(ws.max_row, 26)
        self.assertEqual(ws.max_column, 12)
        self.assertEqual(ws["C1"].value, "Split Model")
        self.assertEqual(ws["F1"].value, "Pooled Model(Non-split)")
        self.assertEqual(ws["A3"].value, "Weather Exclude")
        self.assertEqual(ws["B3"].value, "4-month horizon")


if __name__ == "__main__":
    unittest.main()
